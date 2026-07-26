from __future__ import annotations

import json
import hashlib
import math
import re
import shutil
import sqlite3
import zipfile
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Iterator

from openpyxl.utils import get_column_letter

from .cause_effect import (
    CauseEffectWorkbook,
    normalise_zone_key,
    read_cause_effect_workbook,
)
from .models import Change, ParsedNcf
from .ncf import parse_configuration, parse_ncf


SCHEMA_VERSION = 12


SCHEMA = """
PRAGMA foreign_keys = ON;

CREATE TABLE IF NOT EXISTS metadata (
    key TEXT PRIMARY KEY,
    value TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS snapshots (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    imported_at TEXT NOT NULL,
    source_name TEXT NOT NULL,
    source_path TEXT NOT NULL,
    sha256 TEXT NOT NULL UNIQUE,
    versions_json TEXT NOT NULL,
    warnings_json TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS panels (
    snapshot_id INTEGER NOT NULL REFERENCES snapshots(id) ON DELETE CASCADE,
    node INTEGER NOT NULL,
    name TEXT NOT NULL,
    pcf_bytes INTEGER NOT NULL,
    loops_json TEXT NOT NULL,
    point_table_offset INTEGER,
    point_sub_records INTEGER NOT NULL DEFAULT 0,
    PRIMARY KEY (snapshot_id, node)
);

CREATE TABLE IF NOT EXISTS devices (
    snapshot_id INTEGER NOT NULL REFERENCES snapshots(id) ON DELETE CASCADE,
    stable_key TEXT NOT NULL,
    node INTEGER NOT NULL,
    panel TEXT NOT NULL,
    loop INTEGER NOT NULL,
    address INTEGER NOT NULL,
    sub_address INTEGER NOT NULL,
    zone INTEGER NOT NULL,
    text TEXT NOT NULL,
    product_code INTEGER NOT NULL,
    observed_type TEXT,
    output_group INTEGER,
    output_group_name TEXT,
    ringing_style TEXT,
    record_offset INTEGER NOT NULL,
    PRIMARY KEY (snapshot_id, stable_key)
);

CREATE INDEX IF NOT EXISTS ix_devices_snapshot_zone
    ON devices(snapshot_id, zone);
CREATE INDEX IF NOT EXISTS ix_devices_snapshot_node_loop
    ON devices(snapshot_id, node, loop);

CREATE TABLE IF NOT EXISTS zones (
    number INTEGER PRIMARY KEY,
    description TEXT NOT NULL DEFAULT '',
    floor_id INTEGER REFERENCES floors(id),
    clinical_area INTEGER NOT NULL DEFAULT 1
);

CREATE TABLE IF NOT EXISTS floors (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    level_order INTEGER NOT NULL,
    dxf_path TEXT,
    UNIQUE(name)
);

CREATE TABLE IF NOT EXISTS zone_geometry (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    zone INTEGER NOT NULL REFERENCES zones(number) ON DELETE CASCADE,
    floor_id INTEGER NOT NULL REFERENCES floors(id) ON DELETE CASCADE,
    geometry_json TEXT NOT NULL,
    source_layer TEXT,
    UNIQUE(zone, floor_id)
);

CREATE TABLE IF NOT EXISTS ignored_zone_shapes (
    floor_id INTEGER NOT NULL REFERENCES floors(id) ON DELETE CASCADE,
    shape_key TEXT NOT NULL,
    geometry_json TEXT NOT NULL,
    source_layer TEXT,
    PRIMARY KEY (floor_id, shape_key)
);

CREATE TABLE IF NOT EXISTS map_assets (
    entity_kind TEXT NOT NULL CHECK(entity_kind IN ('device', 'panel')),
    entity_key TEXT NOT NULL,
    floor_id INTEGER NOT NULL REFERENCES floors(id) ON DELETE CASCADE,
    x REAL NOT NULL,
    y REAL NOT NULL,
    symbol_type TEXT NOT NULL,
    PRIMARY KEY (entity_kind, entity_key)
);

CREATE INDEX IF NOT EXISTS ix_map_assets_floor
    ON map_assets(floor_id);

CREATE TABLE IF NOT EXISTS doors (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    floor_id INTEGER NOT NULL REFERENCES floors(id) ON DELETE CASCADE,
    start_x REAL NOT NULL,
    start_y REAL NOT NULL,
    end_x REAL NOT NULL,
    end_y REAL NOT NULL,
    zone_a INTEGER NOT NULL REFERENCES zones(number),
    zone_b INTEGER NOT NULL REFERENCES zones(number),
    has_access_control INTEGER NOT NULL DEFAULT 0,
    access_device_key TEXT,
    access_normal_state TEXT NOT NULL DEFAULT 'LOCKED'
        CHECK(access_normal_state IN ('LOCKED', 'UNLOCKED')),
    has_hold_open INTEGER NOT NULL DEFAULT 0,
    hold_open_device_key TEXT,
    hold_open_normal_state TEXT NOT NULL DEFAULT 'HELD OPEN'
        CHECK(hold_open_normal_state IN ('HELD OPEN', 'CLOSED')),
    door_type TEXT NOT NULL DEFAULT 'SINGLE'
        CHECK(door_type IN ('SINGLE', 'DOUBLE')),
    sprite_x REAL,
    sprite_y REAL,
    rotation_degrees REAL NOT NULL DEFAULT 0,
    notes TEXT NOT NULL DEFAULT '',
    CHECK(has_access_control = 1 OR has_hold_open = 1)
);

CREATE INDEX IF NOT EXISTS ix_doors_floor
    ON doors(floor_id);
CREATE INDEX IF NOT EXISTS ix_doors_zones
    ON doors(zone_a, zone_b);

CREATE TABLE IF NOT EXISTS output_groups (
    number INTEGER PRIMARY KEY,
    name TEXT NOT NULL DEFAULT ''
);

CREATE TABLE IF NOT EXISTS output_group_zone_assignments (
    node INTEGER NOT NULL,
    output_group INTEGER NOT NULL,
    zone INTEGER NOT NULL REFERENCES zones(number) ON DELETE CASCADE,
    output_kind TEXT NOT NULL
        CHECK(output_kind IN ('SOUNDER', 'BEACON')),
    PRIMARY KEY (node, output_group, zone, output_kind)
);

CREATE TABLE IF NOT EXISTS configuration_output_group_lines (
    snapshot_id INTEGER NOT NULL REFERENCES snapshots(id) ON DELETE CASCADE,
    source_row INTEGER NOT NULL,
    target_node INTEGER NOT NULL,
    target_node_name TEXT NOT NULL DEFAULT '',
    output_group INTEGER NOT NULL,
    output_group_name TEXT NOT NULL DEFAULT '',
    operation INTEGER NOT NULL,
    output_style_number INTEGER NOT NULL,
    ringing_style TEXT NOT NULL DEFAULT '',
    ringing_style_name TEXT NOT NULL DEFAULT '',
    zone_from INTEGER NOT NULL,
    zone_to INTEGER NOT NULL,
    zone_qualifiers INTEGER NOT NULL DEFAULT 0,
    PRIMARY KEY (snapshot_id, source_row, target_node, output_group)
);

CREATE INDEX IF NOT EXISTS ix_configuration_output_group_zone_range
    ON configuration_output_group_lines(
        snapshot_id, zone_from, zone_to, target_node
    );

CREATE TABLE IF NOT EXISTS cause_effect_rules (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    trigger_zone INTEGER,
    relation TEXT NOT NULL DEFAULT 'exact',
    target_zone INTEGER,
    target_node INTEGER,
    output_group INTEGER,
    action TEXT NOT NULL,
    source TEXT NOT NULL DEFAULT 'custom',
    enabled INTEGER NOT NULL DEFAULT 1,
    notes TEXT NOT NULL DEFAULT ''
);

CREATE TABLE IF NOT EXISTS cause_effect_imports (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    imported_at TEXT NOT NULL,
    source_name TEXT NOT NULL,
    source_path TEXT NOT NULL,
    sha256 TEXT NOT NULL UNIQUE,
    activation_count INTEGER NOT NULL,
    reference_count INTEGER NOT NULL,
    matched_count INTEGER NOT NULL,
    matrix_only_count INTEGER NOT NULL,
    reference_only_count INTEGER NOT NULL,
    activation_codes_json TEXT NOT NULL,
    warnings_json TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS cause_effect_output_groups (
    import_id INTEGER NOT NULL REFERENCES cause_effect_imports(id) ON DELETE CASCADE,
    target_node INTEGER NOT NULL,
    target_node_name TEXT NOT NULL DEFAULT '',
    output_group INTEGER NOT NULL,
    output_group_name TEXT NOT NULL DEFAULT '',
    source_column TEXT NOT NULL,
    PRIMARY KEY (import_id, target_node, output_group)
);

CREATE TABLE IF NOT EXISTS cause_effect_activations (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    import_id INTEGER NOT NULL REFERENCES cause_effect_imports(id) ON DELETE CASCADE,
    trigger_zone TEXT NOT NULL,
    trigger_zone_name TEXT NOT NULL DEFAULT '',
    target_node INTEGER NOT NULL,
    target_node_name TEXT NOT NULL DEFAULT '',
    output_group INTEGER NOT NULL,
    output_group_name TEXT NOT NULL DEFAULT '',
    output_zone_name TEXT NOT NULL DEFAULT '',
    ringing_style TEXT NOT NULL,
    source_row INTEGER NOT NULL,
    source_column TEXT NOT NULL,
    reference_status TEXT NOT NULL DEFAULT 'matrix_only',
    comments TEXT NOT NULL DEFAULT ''
);

CREATE INDEX IF NOT EXISTS ix_cause_effect_activations_import_zone
    ON cause_effect_activations(import_id, trigger_zone);
CREATE INDEX IF NOT EXISTS ix_cause_effect_activations_import_output
    ON cause_effect_activations(import_id, target_node, output_group);

CREATE TABLE IF NOT EXISTS cause_effect_reference_only (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    import_id INTEGER NOT NULL REFERENCES cause_effect_imports(id) ON DELETE CASCADE,
    trigger_zone TEXT NOT NULL,
    target_node INTEGER NOT NULL,
    target_node_name TEXT NOT NULL DEFAULT '',
    output_group INTEGER NOT NULL,
    output_group_name TEXT NOT NULL DEFAULT '',
    output_zone_name TEXT NOT NULL DEFAULT '',
    ringing_style TEXT NOT NULL,
    source_row INTEGER NOT NULL
);

CREATE TABLE IF NOT EXISTS cause_effect_changes (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    import_id INTEGER NOT NULL REFERENCES cause_effect_imports(id) ON DELETE CASCADE,
    entity TEXT NOT NULL,
    stable_key TEXT NOT NULL,
    change_type TEXT NOT NULL,
    field TEXT,
    old_value TEXT,
    new_value TEXT
);

CREATE INDEX IF NOT EXISTS ix_cause_effect_changes_import
    ON cause_effect_changes(import_id, id);

CREATE TABLE IF NOT EXISTS node_power (
    node INTEGER PRIMARY KEY,
    battery_ah REAL,
    standby_hours REAL NOT NULL DEFAULT 24,
    alarm_minutes REAL NOT NULL DEFAULT 30,
    safety_factor REAL NOT NULL DEFAULT 1.25
);

CREATE TABLE IF NOT EXISTS changes (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    snapshot_id INTEGER NOT NULL REFERENCES snapshots(id) ON DELETE CASCADE,
    entity TEXT NOT NULL,
    stable_key TEXT NOT NULL,
    change_type TEXT NOT NULL,
    field TEXT,
    old_value TEXT,
    new_value TEXT
);

CREATE TABLE IF NOT EXISTS test_sessions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    started_at TEXT NOT NULL,
    completed_at TEXT,
    engineer TEXT NOT NULL DEFAULT '',
    scope_node INTEGER,
    notes TEXT NOT NULL DEFAULT ''
);

CREATE TABLE IF NOT EXISTS test_results (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    session_id INTEGER NOT NULL REFERENCES test_sessions(id) ON DELETE CASCADE,
    trigger_zone INTEGER NOT NULL,
    stable_key TEXT,
    expected_state TEXT NOT NULL,
    actual_state TEXT,
    result TEXT NOT NULL DEFAULT 'not-tested',
    comments TEXT NOT NULL DEFAULT '',
    tested_at TEXT
);

CREATE TABLE IF NOT EXISTS project_audit_log (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    changed_at TEXT NOT NULL,
    table_name TEXT NOT NULL,
    record_key TEXT NOT NULL,
    change_type TEXT NOT NULL
        CHECK(change_type IN ('added', 'modified', 'removed')),
    old_values_json TEXT,
    new_values_json TEXT
);

CREATE INDEX IF NOT EXISTS ix_project_audit_log_changed_at
    ON project_audit_log(changed_at, id);
"""


def zone_shape_key(points: Iterable[tuple[float, float]]) -> str:
    canonical = [
        [round(float(x), 5), round(float(y), 5)]
        for x, y in points
    ]
    payload = json.dumps(
        canonical,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


class ProjectRepository:
    def __init__(self, path: str | Path):
        self.path = Path(path).resolve()
        if not self.path.exists():
            raise FileNotFoundError(self.path)
        self._initialise()

    @classmethod
    def create(cls, path: str | Path, name: str, ncf_path: str | Path) -> "ProjectRepository":
        """Create a project from a legacy NCF or newer SKF configuration."""
        return cls.create_from_source(path, name, ncf_path)

    @classmethod
    def create_from_source(
        cls,
        path: str | Path,
        name: str,
        source_path: str | Path,
    ) -> "ProjectRepository":
        """Create a project from a configuration or Cause & Effect workbook."""
        project_path = Path(path).resolve()
        project_path.parent.mkdir(parents=True, exist_ok=True)
        if project_path.exists():
            raise FileExistsError(project_path)
        sqlite3.connect(project_path).close()
        repository = cls(project_path)
        repository.set_metadata("project_name", name)
        repository.set_metadata("created_at", _now())
        initial_source = Path(source_path)
        if initial_source.suffix.casefold() == ".xlsx":
            repository.import_cause_effect(initial_source)
        elif initial_source.suffix.casefold() in {".ncf", ".skf"}:
            repository.import_configuration(initial_source)
        else:
            raise ValueError(
                "Initial project source must be a Cause & Effect workbook "
                "(.xlsx) or network configuration (.ncf/.skf)."
            )
        return repository

    def _initialise(self) -> None:
        with self.connection() as connection:
            connection.executescript(SCHEMA)
            device_columns = {
                row["name"] for row in connection.execute("PRAGMA table_info(devices)")
            }
            if "output_group_name" not in device_columns:
                connection.execute("ALTER TABLE devices ADD COLUMN output_group_name TEXT")
            if "ringing_style" not in device_columns:
                connection.execute("ALTER TABLE devices ADD COLUMN ringing_style TEXT")
            door_columns = {
                row["name"] for row in connection.execute("PRAGMA table_info(doors)")
            }
            if "door_type" not in door_columns:
                connection.execute(
                    "ALTER TABLE doors ADD COLUMN door_type TEXT NOT NULL DEFAULT 'SINGLE'"
                )
            if "sprite_x" not in door_columns:
                connection.execute("ALTER TABLE doors ADD COLUMN sprite_x REAL")
            if "sprite_y" not in door_columns:
                connection.execute("ALTER TABLE doors ADD COLUMN sprite_y REAL")
            if "rotation_degrees" not in door_columns:
                connection.execute(
                    "ALTER TABLE doors ADD COLUMN rotation_degrees REAL NOT NULL DEFAULT 0"
                )
            self._migrate_internal_zone_doors(connection)
            connection.execute(
                """
                UPDATE doors
                SET sprite_x = (start_x + end_x) / 2.0,
                    sprite_y = (start_y + end_y) / 2.0
                WHERE sprite_x IS NULL OR sprite_y IS NULL
                """
            )
            connection.execute(
                """
                INSERT INTO metadata(key, value)
                VALUES('schema_version', ?)
                ON CONFLICT(key) DO UPDATE SET value = excluded.value
                WHERE metadata.value IS NOT excluded.value
                """,
                (str(SCHEMA_VERSION),),
            )
        self._backfill_zone_descriptions()
        self._backfill_output_groups()
        self._backfill_cause_effect_output_groups()
        self._backfill_configuration_output_group_lines()
        self._backfill_configuration_output_group_changes()
        self._initialise_audit_triggers()

    def _initialise_audit_triggers(self) -> None:
        auditable_tables = (
            "metadata",
            "snapshots",
            "zones",
            "floors",
            "zone_geometry",
            "ignored_zone_shapes",
            "map_assets",
            "doors",
            "output_groups",
            "output_group_zone_assignments",
            "cause_effect_rules",
            "cause_effect_imports",
            "cause_effect_activations",
            "node_power",
            "test_sessions",
            "test_results",
        )
        with self.connection() as connection:
            for table in auditable_tables:
                columns = list(
                    connection.execute(f'PRAGMA table_info("{table}")')
                )
                if not columns:
                    continue
                names = [str(column["name"]) for column in columns]
                primary_keys = [
                    str(column["name"])
                    for column in sorted(
                        columns,
                        key=lambda column: int(column["pk"]),
                    )
                    if int(column["pk"]) > 0
                ]
                key_names = primary_keys or names[:1]

                def json_expression(prefix: str, selected: list[str]) -> str:
                    values = ", ".join(
                        f"'{name}', {prefix}.\"{name}\""
                        for name in selected
                    )
                    return f"json_object({values})"

                new_values = json_expression("NEW", names)
                old_values = json_expression("OLD", names)
                new_key = json_expression("NEW", key_names)
                old_key = json_expression("OLD", key_names)
                changed = " OR ".join(
                    f'OLD."{name}" IS NOT NEW."{name}"'
                    for name in names
                )
                trigger_prefix = f"audit_{table}"
                connection.executescript(
                    f"""
                    CREATE TRIGGER IF NOT EXISTS "{trigger_prefix}_insert"
                    AFTER INSERT ON "{table}"
                    BEGIN
                        INSERT INTO project_audit_log(
                            changed_at, table_name, record_key, change_type,
                            old_values_json, new_values_json
                        ) VALUES (
                            strftime('%Y-%m-%dT%H:%M:%fZ', 'now'),
                            '{table}', {new_key}, 'added', NULL, {new_values}
                        );
                    END;

                    CREATE TRIGGER IF NOT EXISTS "{trigger_prefix}_update"
                    AFTER UPDATE ON "{table}"
                    WHEN {changed}
                    BEGIN
                        INSERT INTO project_audit_log(
                            changed_at, table_name, record_key, change_type,
                            old_values_json, new_values_json
                        ) VALUES (
                            strftime('%Y-%m-%dT%H:%M:%fZ', 'now'),
                            '{table}', {new_key}, 'modified',
                            {old_values}, {new_values}
                        );
                    END;

                    CREATE TRIGGER IF NOT EXISTS "{trigger_prefix}_delete"
                    AFTER DELETE ON "{table}"
                    BEGIN
                        INSERT INTO project_audit_log(
                            changed_at, table_name, record_key, change_type,
                            old_values_json, new_values_json
                        ) VALUES (
                            strftime('%Y-%m-%dT%H:%M:%fZ', 'now'),
                            '{table}', {old_key}, 'removed', {old_values}, NULL
                        );
                    END;
                    """
                )

    @staticmethod
    def _migrate_internal_zone_doors(connection: sqlite3.Connection) -> None:
        schema_row = connection.execute(
            "SELECT sql FROM sqlite_master WHERE type = 'table' AND name = 'doors'"
        ).fetchone()
        normalised_schema = re.sub(
            r"\s+",
            "",
            str(schema_row["sql"] if schema_row else ""),
        ).upper()
        if "CHECK(ZONE_A<>ZONE_B)" not in normalised_schema:
            return
        connection.execute("DROP INDEX IF EXISTS ix_doors_floor")
        connection.execute("DROP INDEX IF EXISTS ix_doors_zones")
        connection.execute("ALTER TABLE doors RENAME TO doors_distinct_zone_legacy")
        connection.execute(
            """
            CREATE TABLE doors (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                floor_id INTEGER NOT NULL REFERENCES floors(id) ON DELETE CASCADE,
                start_x REAL NOT NULL,
                start_y REAL NOT NULL,
                end_x REAL NOT NULL,
                end_y REAL NOT NULL,
                zone_a INTEGER NOT NULL REFERENCES zones(number),
                zone_b INTEGER NOT NULL REFERENCES zones(number),
                has_access_control INTEGER NOT NULL DEFAULT 0,
                access_device_key TEXT,
                access_normal_state TEXT NOT NULL DEFAULT 'LOCKED'
                    CHECK(access_normal_state IN ('LOCKED', 'UNLOCKED')),
                has_hold_open INTEGER NOT NULL DEFAULT 0,
                hold_open_device_key TEXT,
                hold_open_normal_state TEXT NOT NULL DEFAULT 'HELD OPEN'
                    CHECK(hold_open_normal_state IN ('HELD OPEN', 'CLOSED')),
                door_type TEXT NOT NULL DEFAULT 'SINGLE'
                    CHECK(door_type IN ('SINGLE', 'DOUBLE')),
                sprite_x REAL,
                sprite_y REAL,
                rotation_degrees REAL NOT NULL DEFAULT 0,
                notes TEXT NOT NULL DEFAULT '',
                CHECK(has_access_control = 1 OR has_hold_open = 1)
            )
            """
        )
        connection.execute(
            """
            INSERT INTO doors(
                id, name, floor_id, start_x, start_y, end_x, end_y,
                zone_a, zone_b, has_access_control, access_device_key,
                access_normal_state, has_hold_open, hold_open_device_key,
                hold_open_normal_state, door_type, sprite_x, sprite_y,
                rotation_degrees, notes
            )
            SELECT
                id, name, floor_id, start_x, start_y, end_x, end_y,
                zone_a, zone_b, has_access_control, access_device_key,
                access_normal_state, has_hold_open, hold_open_device_key,
                hold_open_normal_state, door_type, sprite_x, sprite_y,
                rotation_degrees, notes
            FROM doors_distinct_zone_legacy
            """
        )
        connection.execute("DROP TABLE doors_distinct_zone_legacy")
        connection.execute(
            "CREATE INDEX ix_doors_floor ON doors(floor_id)"
        )
        connection.execute(
            "CREATE INDEX ix_doors_zones ON doors(zone_a, zone_b)"
        )

    def _backfill_zone_descriptions(self) -> None:
        """Populate projects created before SITE zone names were decoded."""
        with self.connection() as connection:
            counts = connection.execute(
                """
                SELECT COUNT(*) AS total,
                       SUM(CASE WHEN TRIM(description) <> '' THEN 1 ELSE 0 END) AS named
                FROM zones
                """
            ).fetchone()
            snapshot = connection.execute(
                "SELECT source_path FROM snapshots ORDER BY id DESC LIMIT 1"
            ).fetchone()
        if (
            not counts
            or not counts["total"]
            or counts["named"]
            or not snapshot
            or not Path(snapshot["source_path"]).exists()
        ):
            return
        try:
            parsed = parse_ncf(snapshot["source_path"])
        except (OSError, ValueError, zipfile.BadZipFile):
            return
        with self.connection() as connection:
            self._upsert_zones(connection, parsed)

    def _backfill_output_groups(self) -> None:
        """Populate native group data in projects imported by older parsers."""
        with self.connection() as connection:
            snapshot = connection.execute(
                "SELECT id, source_path FROM snapshots ORDER BY id DESC LIMIT 1"
            ).fetchone()
            populated = connection.execute(
                "SELECT 1 FROM devices WHERE output_group IS NOT NULL LIMIT 1"
            ).fetchone()
        if (
            populated
            or not snapshot
            or not Path(snapshot["source_path"]).exists()
        ):
            return
        try:
            parsed = parse_ncf(snapshot["source_path"])
        except (OSError, ValueError, zipfile.BadZipFile):
            return
        rows = []
        for panel in parsed.panels:
            for device in panel.devices:
                for channel in device.channels:
                    rows.append(
                        (
                            channel.output_group,
                            channel.output_group_name,
                            channel.ringing_style,
                            snapshot["id"],
                            f"{device.node}/{device.loop}/{device.address}/{channel.sub_address}",
                        )
                    )
        with self.connection() as connection:
            connection.executemany(
                """
                UPDATE devices
                SET output_group = ?, output_group_name = ?, ringing_style = ?
                WHERE snapshot_id = ? AND stable_key = ?
                """,
                rows,
            )

    def _backfill_cause_effect_output_groups(self) -> None:
        """Retain empty matrix output columns in projects imported before v6."""
        with self.connection() as connection:
            latest = connection.execute(
                """
                SELECT id, source_path
                FROM cause_effect_imports
                ORDER BY id DESC
                LIMIT 1
                """
            ).fetchone()
            if latest is None:
                return
            populated = connection.execute(
                """
                SELECT 1
                FROM cause_effect_output_groups
                WHERE import_id = ?
                LIMIT 1
                """,
                (latest["id"],),
            ).fetchone()
        source = Path(latest["source_path"])
        if populated or not source.exists():
            return
        try:
            parsed = read_cause_effect_workbook(source)
        except (OSError, ValueError, zipfile.BadZipFile):
            return
        with self.connection() as connection:
            self._insert_cause_effect_output_groups(
                connection,
                int(latest["id"]),
                parsed,
            )

    def _backfill_configuration_output_group_lines(self) -> None:
        with self.connection() as connection:
            snapshot = connection.execute(
                """
                SELECT id, source_path
                FROM snapshots
                ORDER BY id DESC
                LIMIT 1
                """
            ).fetchone()
            if snapshot is None:
                return
            populated = connection.execute(
                """
                SELECT 1
                FROM configuration_output_group_lines
                WHERE snapshot_id = ?
                LIMIT 1
                """,
                (snapshot["id"],),
            ).fetchone()
        source = Path(snapshot["source_path"])
        if populated or not source.exists():
            return
        try:
            parsed = parse_configuration(source)
        except (OSError, ValueError, zipfile.BadZipFile):
            return
        if not parsed.output_group_lines:
            return
        with self.connection() as connection:
            connection.executemany(
                """
                INSERT OR REPLACE INTO configuration_output_group_lines(
                    snapshot_id, source_row, target_node, target_node_name,
                    output_group, output_group_name, operation,
                    output_style_number, ringing_style, ringing_style_name,
                    zone_from, zone_to, zone_qualifiers
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    (
                        int(snapshot["id"]),
                        line.source_row,
                        line.target_node,
                        line.target_node_name,
                        line.output_group,
                        line.output_group_name,
                        line.operation,
                        line.output_style_number,
                        line.ringing_style,
                        line.ringing_style_name,
                        line.zone_from,
                        line.zone_to,
                        line.zone_qualifiers,
                    )
                    for line in parsed.output_group_lines
                ],
            )

    def _backfill_configuration_output_group_changes(self) -> None:
        with self.connection() as connection:
            snapshot_ids = [
                int(row["snapshot_id"])
                for row in connection.execute(
                    """
                    SELECT DISTINCT snapshot_id
                    FROM configuration_output_group_lines
                    ORDER BY snapshot_id
                    """
                )
            ]
            previous_snapshot = None
            for snapshot_id in snapshot_ids:
                if previous_snapshot is None:
                    count = connection.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM (
                            SELECT target_node, output_group
                            FROM configuration_output_group_lines
                            WHERE snapshot_id = ?
                            GROUP BY target_node, output_group
                        )
                        """,
                        (snapshot_id,),
                    ).fetchone()["count"]
                    generated = [
                        Change(
                            "snapshot",
                            str(snapshot_id),
                            "initial",
                            "output_groups",
                            None,
                            str(count),
                        )
                    ]
                else:
                    generated = [
                        change
                        for change in self._calculate_changes(
                            connection,
                            previous_snapshot,
                            snapshot_id,
                        )
                        if change.entity == "output_group"
                    ]
                existing = {
                    (
                        row["entity"],
                        row["stable_key"],
                        row["change_type"],
                        row["field"],
                        row["old_value"],
                        row["new_value"],
                    )
                    for row in connection.execute(
                        """
                        SELECT entity, stable_key, change_type, field,
                               old_value, new_value
                        FROM changes
                        WHERE snapshot_id = ?
                        """,
                        (snapshot_id,),
                    )
                }
                missing = [
                    change
                    for change in generated
                    if (
                        change.entity,
                        change.stable_key,
                        change.change_type,
                        change.field,
                        change.old_value,
                        change.new_value,
                    )
                    not in existing
                ]
                connection.executemany(
                    """
                    INSERT INTO changes(
                        snapshot_id, entity, stable_key, change_type,
                        field, old_value, new_value
                    ) VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            snapshot_id,
                            change.entity,
                            change.stable_key,
                            change.change_type,
                            change.field,
                            change.old_value,
                            change.new_value,
                        )
                        for change in missing
                    ],
                )
                previous_snapshot = snapshot_id

    @contextmanager
    def connection(self) -> Iterator[sqlite3.Connection]:
        connection = sqlite3.connect(self.path)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        try:
            yield connection
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def metadata(self, key: str, default: str = "") -> str:
        with self.connection() as connection:
            row = connection.execute("SELECT value FROM metadata WHERE key = ?", (key,)).fetchone()
        return row["value"] if row else default

    def set_metadata(self, key: str, value: str) -> None:
        with self.connection() as connection:
            connection.execute(
                "INSERT OR REPLACE INTO metadata(key, value) VALUES(?, ?)",
                (key, value),
            )

    @property
    def name(self) -> str:
        return self.metadata("project_name", self.path.stem)

    def save_as(self, target: str | Path) -> "ProjectRepository":
        target_path = Path(target).resolve()
        if target_path.exists():
            raise FileExistsError(target_path)
        with self.connection() as connection:
            connection.execute("PRAGMA wal_checkpoint(FULL)")
        shutil.copy2(self.path, target_path)
        return ProjectRepository(target_path)

    def latest_snapshot_id(self) -> int | None:
        with self.connection() as connection:
            row = connection.execute("SELECT MAX(id) AS id FROM snapshots").fetchone()
        return int(row["id"]) if row and row["id"] is not None else None

    def import_configuration(
        self, configuration_path: str | Path
    ) -> tuple[int, list[Change]]:
        parsed = parse_configuration(configuration_path)
        with self.connection() as connection:
            existing = connection.execute(
                "SELECT id FROM snapshots WHERE sha256 = ?", (parsed.sha256,)
            ).fetchone()
            if existing:
                self._upsert_zones(connection, parsed)
                return int(existing["id"]), []

            previous_snapshot = connection.execute("SELECT MAX(id) AS id FROM snapshots").fetchone()["id"]
            cursor = connection.execute(
                """
                INSERT INTO snapshots(
                    imported_at, source_name, source_path, sha256, versions_json, warnings_json
                ) VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    _now(),
                    parsed.source.name,
                    str(parsed.source),
                    parsed.sha256,
                    json.dumps(parsed.versions),
                    json.dumps(parsed.warnings),
                ),
            )
            snapshot_id = int(cursor.lastrowid)
            self._insert_parsed(connection, snapshot_id, parsed)
            changes = self._calculate_changes(connection, previous_snapshot, snapshot_id)
            connection.executemany(
                """
                INSERT INTO changes(
                    snapshot_id, entity, stable_key, change_type, field, old_value, new_value
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    (
                        snapshot_id,
                        change.entity,
                        change.stable_key,
                        change.change_type,
                        change.field,
                        change.old_value,
                        change.new_value,
                    )
                    for change in changes
                ],
            )
        return snapshot_id, changes

    def import_ncf(self, ncf_path: str | Path) -> tuple[int, list[Change]]:
        """Backward-compatible alias for legacy callers."""
        return self.import_configuration(ncf_path)

    @staticmethod
    def _insert_parsed(
        connection: sqlite3.Connection, snapshot_id: int, parsed: ParsedNcf
    ) -> None:
        connection.executemany(
            """
            INSERT INTO panels(
                snapshot_id, node, name, pcf_bytes, loops_json,
                point_table_offset, point_sub_records
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    snapshot_id,
                    panel.node,
                    panel.name,
                    panel.pcf_bytes,
                    json.dumps(panel.loops),
                    panel.point_table_offset,
                    panel.point_sub_records,
                )
                for panel in parsed.panels
            ],
        )

        device_rows = []
        for panel in parsed.panels:
            for device in panel.devices:
                for channel in device.channels:
                    stable_key = f"{device.node}/{device.loop}/{device.address}/{channel.sub_address}"
                    device_rows.append(
                        (
                            snapshot_id,
                            stable_key,
                            device.node,
                            device.panel,
                            device.loop,
                            device.address,
                            channel.sub_address,
                            channel.zone,
                            channel.text,
                            channel.product_code,
                            channel.observed_type,
                            channel.output_group,
                            channel.output_group_name,
                            channel.ringing_style,
                            channel.record_offset,
                        )
                    )
        connection.executemany(
            """
            INSERT INTO devices(
                snapshot_id, stable_key, node, panel, loop, address, sub_address,
                zone, text, product_code, observed_type, output_group,
                output_group_name, ringing_style, record_offset
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            device_rows,
        )
        connection.executemany(
            """
            INSERT INTO configuration_output_group_lines(
                snapshot_id, source_row, target_node, target_node_name,
                output_group, output_group_name, operation,
                output_style_number, ringing_style, ringing_style_name,
                zone_from, zone_to, zone_qualifiers
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    snapshot_id,
                    line.source_row,
                    line.target_node,
                    line.target_node_name,
                    line.output_group,
                    line.output_group_name,
                    line.operation,
                    line.output_style_number,
                    line.ringing_style,
                    line.ringing_style_name,
                    line.zone_from,
                    line.zone_to,
                    line.zone_qualifiers,
                )
                for line in parsed.output_group_lines
            ],
        )
        ProjectRepository._upsert_zones(connection, parsed)

    @staticmethod
    def _upsert_zones(
        connection: sqlite3.Connection, parsed: ParsedNcf
    ) -> None:
        connection.executemany(
            """
            INSERT INTO zones(number, description) VALUES (?, ?)
            ON CONFLICT(number) DO UPDATE SET
                description = CASE
                    WHEN TRIM(zones.description) = '' THEN excluded.description
                    ELSE zones.description
                END
            """,
            [(zone.number, zone.description) for zone in parsed.zones],
        )

    @staticmethod
    def _calculate_changes(
        connection: sqlite3.Connection, previous_snapshot: int | None, snapshot_id: int
    ) -> list[Change]:
        if previous_snapshot is None:
            count = connection.execute(
                "SELECT COUNT(*) AS count FROM devices WHERE snapshot_id = ?", (snapshot_id,)
            ).fetchone()["count"]
            output_group_count = connection.execute(
                """
                SELECT COUNT(*) AS count
                FROM (
                    SELECT target_node, output_group
                    FROM configuration_output_group_lines
                    WHERE snapshot_id = ?
                    GROUP BY target_node, output_group
                )
                """,
                (snapshot_id,),
            ).fetchone()["count"]
            return [
                Change(
                    entity="snapshot",
                    stable_key=str(snapshot_id),
                    change_type="initial",
                    field="devices",
                    old_value=None,
                    new_value=str(count),
                ),
                Change(
                    entity="snapshot",
                    stable_key=str(snapshot_id),
                    change_type="initial",
                    field="output_groups",
                    old_value=None,
                    new_value=str(output_group_count),
                ),
            ]

        fields = (
            "panel",
            "zone",
            "text",
            "product_code",
            "observed_type",
            "output_group",
            "output_group_name",
            "ringing_style",
        )
        old_rows = {
            row["stable_key"]: row
            for row in connection.execute(
                "SELECT * FROM devices WHERE snapshot_id = ?", (previous_snapshot,)
            )
        }
        new_rows = {
            row["stable_key"]: row
            for row in connection.execute(
                "SELECT * FROM devices WHERE snapshot_id = ?", (snapshot_id,)
            )
        }
        changes: list[Change] = []
        for key in sorted(new_rows.keys() - old_rows.keys()):
            changes.append(Change("device", key, "added", None, None, new_rows[key]["text"]))
        for key in sorted(old_rows.keys() - new_rows.keys()):
            changes.append(Change("device", key, "removed", None, old_rows[key]["text"], None))
        for key in sorted(old_rows.keys() & new_rows.keys()):
            for field in fields:
                old_value = old_rows[key][field]
                new_value = new_rows[key][field]
                if old_value != new_value:
                    changes.append(
                        Change(
                            "device",
                            key,
                            "modified",
                            field,
                            _stringify(old_value),
                            _stringify(new_value),
                        )
                    )
        old_output_groups = _configuration_output_group_values(
            connection, previous_snapshot
        )
        new_output_groups = _configuration_output_group_values(
            connection, snapshot_id
        )
        for key in sorted(new_output_groups.keys() - old_output_groups.keys()):
            changes.append(
                Change(
                    "output_group",
                    key,
                    "added",
                    None,
                    None,
                    _output_group_change_summary(new_output_groups[key]),
                )
            )
        for key in sorted(old_output_groups.keys() - new_output_groups.keys()):
            changes.append(
                Change(
                    "output_group",
                    key,
                    "removed",
                    None,
                    _output_group_change_summary(old_output_groups[key]),
                    None,
                )
            )
        output_group_fields = (
            "target_node_name",
            "output_group_name",
            "zone_triggers",
            "ringing_styles",
            "operations",
            "zone_qualifiers",
        )
        for key in sorted(old_output_groups.keys() & new_output_groups.keys()):
            old_group = old_output_groups[key]
            new_group = new_output_groups[key]
            for field in output_group_fields:
                if old_group[field] != new_group[field]:
                    changes.append(
                        Change(
                            "output_group",
                            key,
                            "modified",
                            field,
                            _stringify(old_group[field]),
                            _stringify(new_group[field]),
                        )
                    )
        return changes

    def fetch_devices(self, snapshot_id: int | None = None) -> list[sqlite3.Row]:
        snapshot_id = snapshot_id or self.latest_snapshot_id()
        if snapshot_id is None:
            return []
        with self.connection() as connection:
            return list(
                connection.execute(
                    """
                    SELECT * FROM devices
                    WHERE snapshot_id = ?
                    ORDER BY node, loop, address, sub_address
                    """,
                    (snapshot_id,),
                )
            )

    def fetch_panels(self, snapshot_id: int | None = None) -> list[sqlite3.Row]:
        snapshot_id = snapshot_id or self.latest_snapshot_id()
        if snapshot_id is None:
            return []
        with self.connection() as connection:
            return list(
                connection.execute(
                    """
                    SELECT p.*,
                           COUNT(DISTINCT d.loop || '/' || d.address) AS device_count
                    FROM panels p
                    LEFT JOIN devices d
                      ON d.snapshot_id = p.snapshot_id AND d.node = p.node
                    WHERE p.snapshot_id = ?
                    GROUP BY p.snapshot_id, p.node
                    ORDER BY p.node
                    """,
                    (snapshot_id,),
                )
            )

    def fetch_output_groups(self, snapshot_id: int | None = None) -> list[dict]:
        snapshot_id = snapshot_id or self.latest_snapshot_id()
        groups: dict[tuple[int, int], dict] = {}
        with self.connection() as connection:
            if snapshot_id is not None:
                for row in connection.execute(
                    """
                    SELECT node, panel, output_group,
                           MAX(COALESCE(output_group_name, '')) AS group_name,
                           COUNT(DISTINCT loop || '/' || address) AS device_count,
                           GROUP_CONCAT(DISTINCT ringing_style) AS ringing_styles
                    FROM devices
                    WHERE snapshot_id = ?
                      AND output_group IS NOT NULL
                      AND output_group > 0
                    GROUP BY node, panel, output_group
                    ORDER BY node, output_group
                    """,
                    (snapshot_id,),
                ):
                    groups[(int(row["node"]), int(row["output_group"]))] = dict(
                        row
                    )
                for row in connection.execute(
                    """
                    SELECT target_node AS node,
                           MAX(target_node_name) AS panel,
                           output_group,
                           MAX(output_group_name) AS group_name,
                           GROUP_CONCAT(DISTINCT ringing_style_name)
                               AS ringing_styles
                    FROM configuration_output_group_lines
                    WHERE snapshot_id = ?
                    GROUP BY target_node, output_group
                    ORDER BY target_node, output_group
                    """,
                    (snapshot_id,),
                ):
                    key = (int(row["node"]), int(row["output_group"]))
                    existing = groups.get(key)
                    if existing is None:
                        groups[key] = {
                            **dict(row),
                            "device_count": 0,
                        }
                    else:
                        if str(row["group_name"] or "").strip():
                            existing["group_name"] = row["group_name"]
                        if str(row["ringing_styles"] or "").strip():
                            existing["ringing_styles"] = row["ringing_styles"]
            latest_import = connection.execute(
                "SELECT MAX(id) AS id FROM cause_effect_imports"
            ).fetchone()["id"]
            if latest_import is not None:
                for row in connection.execute(
                    """
                    SELECT target_node AS node,
                           target_node_name AS panel,
                           output_group,
                           output_group_name AS group_name
                    FROM cause_effect_output_groups
                    WHERE import_id = ?
                    ORDER BY target_node, output_group
                    """,
                    (latest_import,),
                ):
                    key = (int(row["node"]), int(row["output_group"]))
                    existing = groups.get(key)
                    if existing is None:
                        groups[key] = {
                            **dict(row),
                            "device_count": 0,
                            "ringing_styles": None,
                        }
                    else:
                        if str(row["panel"] or "").strip():
                            existing["panel"] = row["panel"]
                        if str(row["group_name"] or "").strip():
                            existing["group_name"] = row["group_name"]
        return [
            groups[key]
            for key in sorted(groups)
        ]

    def fetch_output_group_zone_assignments(
        self,
        node: int | None = None,
        output_group: int | None = None,
    ) -> list[sqlite3.Row]:
        clauses = []
        values: list[int] = []
        if node is not None:
            clauses.append("a.node = ?")
            values.append(int(node))
        if output_group is not None:
            clauses.append("a.output_group = ?")
            values.append(int(output_group))
        where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
        with self.connection() as connection:
            return list(
                connection.execute(
                    f"""
                    SELECT a.*, z.description
                    FROM output_group_zone_assignments a
                    JOIN zones z ON z.number = a.zone
                    {where}
                    ORDER BY a.node, a.output_group, a.zone, a.output_kind
                    """,
                    values,
                )
            )

    def replace_output_group_zone_assignments(
        self,
        node: int,
        output_group: int,
        assignments: Iterable[tuple[int, str]],
    ) -> None:
        normalised = sorted(
            {
                (int(zone), str(output_kind).strip().upper())
                for zone, output_kind in assignments
            }
        )
        invalid = [
            output_kind
            for _zone, output_kind in normalised
            if output_kind not in {"SOUNDER", "BEACON"}
        ]
        if invalid:
            raise ValueError(f"Unsupported output type: {invalid[0]}")
        with self.connection() as connection:
            connection.execute(
                """
                DELETE FROM output_group_zone_assignments
                WHERE node = ? AND output_group = ?
                """,
                (int(node), int(output_group)),
            )
            connection.executemany(
                """
                INSERT INTO output_group_zone_assignments(
                    node, output_group, zone, output_kind
                ) VALUES (?, ?, ?, ?)
                """,
                [
                    (int(node), int(output_group), zone, output_kind)
                    for zone, output_kind in normalised
                ],
            )

    def fetch_output_group_devices(
        self, node: int, output_group: int, snapshot_id: int | None = None
    ) -> list[sqlite3.Row]:
        snapshot_id = snapshot_id or self.latest_snapshot_id()
        if snapshot_id is None:
            return []
        with self.connection() as connection:
            return list(
                connection.execute(
                    """
                    SELECT * FROM devices
                    WHERE snapshot_id = ? AND node = ? AND output_group = ?
                    ORDER BY loop, address, sub_address
                    """,
                    (snapshot_id, node, output_group),
                )
            )

    def fetch_zones(self) -> list[sqlite3.Row]:
        snapshot_id = self.latest_snapshot_id()
        with self.connection() as connection:
            return list(
                connection.execute(
                    """
                    SELECT z.*,
                           (
                               SELECT GROUP_CONCAT(zone_floors.name, ', ')
                               FROM (
                                   SELECT f2.name
                                   FROM zone_geometry g2
                                   JOIN floors f2 ON f2.id = g2.floor_id
                                   WHERE g2.zone = z.number
                                   ORDER BY f2.level_order DESC, f2.name
                               ) AS zone_floors
                           ) AS floor_name,
                           (
                               SELECT MAX(f2.level_order)
                               FROM zone_geometry g2
                               JOIN floors f2 ON f2.id = g2.floor_id
                               WHERE g2.zone = z.number
                           ) AS level_order,
                           COUNT(DISTINCT d.node || '/' || d.loop || '/' || d.address) AS device_count,
                           (
                               SELECT GROUP_CONCAT(zone_nodes.node, ', ')
                               FROM (
                                   SELECT DISTINCT node
                                   FROM devices
                                   WHERE snapshot_id = ? AND zone = z.number
                                   ORDER BY node
                               ) AS zone_nodes
                           ) AS nodes
                    FROM zones z
                    LEFT JOIN devices d ON d.snapshot_id = ? AND d.zone = z.number
                    GROUP BY z.number
                    ORDER BY z.number
                    """,
                    (snapshot_id, snapshot_id),
                )
            )

    def fetch_changes(self, snapshot_id: int | None = None) -> list[sqlite3.Row]:
        requested_snapshot = snapshot_id
        snapshot_id = snapshot_id or self.latest_snapshot_id()
        with self.connection() as connection:
            changes = (
                list(
                    connection.execute(
                        "SELECT * FROM changes WHERE snapshot_id = ? ORDER BY id",
                        (snapshot_id,),
                    )
                )
                if snapshot_id is not None
                else []
            )
            if requested_snapshot is not None:
                return changes
            latest_import = connection.execute(
                "SELECT MAX(id) AS id FROM cause_effect_imports"
            ).fetchone()["id"]
            if latest_import is not None:
                changes.extend(
                    connection.execute(
                        """
                        SELECT id, NULL AS snapshot_id, import_id,
                               entity, stable_key, change_type,
                               field, old_value, new_value
                        FROM cause_effect_changes
                        WHERE import_id = ?
                        ORDER BY id
                        """,
                        (latest_import,),
                    )
                )
            return changes

    def fetch_change_details(
        self,
        snapshot_id: int | None = None,
        include_all_imports: bool = False,
    ) -> list[dict]:
        field_labels = {
            "panel": "Panel name",
            "zone": "Zone",
            "text": "Device text",
            "location": "Device text",
            "location_text": "Device text",
            "location text": "Device text",
            "product_code": "Product code",
            "observed_type": "Device type",
            "output_group": "Output group",
            "output_group_name": "Output group name",
            "ringing_style": "Ringing style",
            "output_groups": "Output groups",
            "target_node_name": "Node name",
            "zone_triggers": "Zone trigger extent",
            "ringing_styles": "Ringing styles",
            "operations": "Operations",
            "zone_qualifiers": "Zone qualifiers",
        }
        details = []
        with self.connection() as connection:
            if include_all_imports:
                source_changes = list(
                    connection.execute(
                        """
                        SELECT *
                        FROM (
                            SELECT c.id, c.snapshot_id, NULL AS import_id,
                                   c.entity, c.stable_key, c.change_type,
                                   c.field, c.old_value, c.new_value,
                                   s.imported_at AS changed_at,
                                   s.source_name,
                                   'Configuration' AS source_kind
                            FROM changes c
                            JOIN snapshots s ON s.id = c.snapshot_id
                            UNION ALL
                            SELECT c.id, NULL AS snapshot_id, c.import_id,
                                   c.entity, c.stable_key, c.change_type,
                                   c.field, c.old_value, c.new_value,
                                   i.imported_at AS changed_at,
                                   i.source_name,
                                   'Cause & Effect' AS source_kind
                            FROM cause_effect_changes c
                            JOIN cause_effect_imports i ON i.id = c.import_id
                        )
                        ORDER BY changed_at, source_kind, id
                        """
                    )
                )
            else:
                source_changes = self.fetch_changes(snapshot_id)
            for source_change in source_changes:
                change = dict(source_change)
                context = None
                if (
                    change["entity"] == "device"
                    and change.get("snapshot_id") is not None
                ):
                    snapshot_id = int(change["snapshot_id"])
                    previous = connection.execute(
                        "SELECT MAX(id) AS id FROM snapshots WHERE id < ?",
                        (snapshot_id,),
                    ).fetchone()["id"]
                    context = connection.execute(
                        """
                        SELECT * FROM devices
                        WHERE snapshot_id = ? AND stable_key = ?
                        """,
                        (snapshot_id, change["stable_key"]),
                    ).fetchone()
                    if context is None and previous is not None:
                        context = connection.execute(
                            """
                            SELECT * FROM devices
                            WHERE snapshot_id = ? AND stable_key = ?
                            """,
                            (previous, change["stable_key"]),
                        ).fetchone()
                    context = dict(context) if context is not None else None
                elif (
                    change["entity"] == "output_group"
                    and change.get("snapshot_id") is not None
                ):
                    snapshot_id = int(change["snapshot_id"])
                    previous = connection.execute(
                        "SELECT MAX(id) AS id FROM snapshots WHERE id < ?",
                        (snapshot_id,),
                    ).fetchone()["id"]
                    group_values = _configuration_output_group_values(
                        connection, snapshot_id
                    )
                    context = group_values.get(change["stable_key"])
                    if context is None and previous is not None:
                        context = _configuration_output_group_values(
                            connection, int(previous)
                        ).get(change["stable_key"])
                raw_field = str(change.get("field") or "").strip()
                field = field_labels.get(
                    raw_field.casefold(),
                    raw_field.replace("_", " ").title(),
                )
                if change["entity"] == "device":
                    if change["change_type"] == "added":
                        description = "Device added"
                    elif change["change_type"] == "removed":
                        description = "Device removed"
                    else:
                        description = f"{field or 'Device'} changed"
                elif change["entity"].startswith("cause_effect"):
                    description = (
                        f"Cause & Effect {field} "
                        f"{change['change_type']}"
                    ).strip()
                elif change["entity"] == "output_group":
                    if change["change_type"] == "added":
                        description = "Output group added"
                    elif change["change_type"] == "removed":
                        description = "Output group removed"
                    else:
                        description = (
                            f"{field or 'Output group'} changed"
                        )
                else:
                    description = (
                        f"{str(change['entity']).replace('_', ' ').title()} "
                        f"{change['change_type']}"
                    )
                key_match = re.search(
                    r"\bnode\s+(\d+)\s*/group\s+(\d+)\b",
                    str(change.get("stable_key") or ""),
                    flags=re.IGNORECASE,
                )
                node = (
                    context.get("node")
                    if context is not None
                    else int(key_match.group(1)) if key_match else None
                )
                output_group = (
                    context.get("output_group")
                    if context is not None
                    else int(key_match.group(2)) if key_match else None
                )
                row = {
                    **change,
                    "changed_at": change.get("changed_at", ""),
                    "source_name": change.get("source_name", ""),
                    "source_kind": change.get("source_kind", ""),
                    "raw_field": raw_field,
                    "field": field,
                    "description": description,
                    "node": node,
                    "panel": (
                        context.get(
                            "panel",
                            context.get("target_node_name", ""),
                        )
                        if context is not None
                        else ""
                    ),
                    "zone": context.get("zone") if context is not None else None,
                    "loop": context.get("loop") if context is not None else None,
                    "address": (
                        context.get("address") if context is not None else None
                    ),
                    "sub_address": (
                        context.get("sub_address")
                        if context is not None
                        else None
                    ),
                    "device_text": (
                        context.get("text", "") if context is not None else ""
                    ),
                    "device_type": (
                        context.get("observed_type", "")
                        if context is not None
                        else ""
                    ),
                    "output_group": output_group,
                    "output_group_name": (
                        context.get("output_group_name", "")
                        if context is not None
                        else ""
                    ),
                    "ringing_style": (
                        context.get(
                            "ringing_style",
                            context.get("ringing_styles", ""),
                        )
                        if context is not None
                        else ""
                    ),
                }
                details.append(row)
        return details

    def fetch_project_history(self) -> list[dict]:
        area_labels = {
            "metadata": "Project details",
            "snapshots": "Configuration import",
            "zones": "Zones",
            "floors": "Floors",
            "zone_geometry": "Zone drawings",
            "ignored_zone_shapes": "Suppressed drawing polygons",
            "map_assets": "Device placements",
            "doors": "Doors",
            "output_groups": "Output groups",
            "output_group_zone_assignments": "Output-to-zone assignments",
            "cause_effect_rules": "Cause & Effect rules",
            "cause_effect_imports": "Cause & Effect imports",
            "cause_effect_activations": "Cause & Effect activations",
            "node_power": "Node power settings",
            "test_sessions": "Test sessions",
            "test_results": "Test results",
        }
        with self.connection() as connection:
            rows = list(
                connection.execute(
                    """
                    SELECT *
                    FROM project_audit_log
                    WHERE NOT (
                        table_name = 'metadata'
                        AND COALESCE(
                            json_extract(new_values_json, '$.key'),
                            json_extract(old_values_json, '$.key')
                        ) = 'schema_version'
                    )
                    ORDER BY id DESC
                    """
                )
            )
        history = []
        for row in rows:
            old_values = json.loads(row["old_values_json"] or "{}")
            new_values = json.loads(row["new_values_json"] or "{}")
            values = new_values or old_values
            node = None
            for name in ("node", "target_node", "scope_node"):
                if values.get(name) not in (None, ""):
                    try:
                        node = int(values[name])
                    except (TypeError, ValueError):
                        pass
                    break
            if node is None:
                entity_key = str(values.get("entity_key") or "")
                match = re.match(r"\s*(\d+)\s*/", entity_key)
                if match:
                    node = int(match.group(1))
            changed_fields = []
            if row["change_type"] == "modified":
                changed_fields = [
                    name
                    for name in sorted(set(old_values) | set(new_values))
                    if old_values.get(name) != new_values.get(name)
                ]
                summary = "; ".join(
                    (
                        f"{name.replace('_', ' ').title()}: "
                        f"{_stringify(old_values.get(name)) or 'blank'} → "
                        f"{_stringify(new_values.get(name)) or 'blank'}"
                    )
                    for name in changed_fields
                )
            else:
                values = (
                    new_values
                    if row["change_type"] == "added"
                    else old_values
                )
                useful = [
                    name
                    for name in (
                        "name",
                        "number",
                        "zone",
                        "floor_id",
                        "entity_key",
                        "output_group",
                        "output_kind",
                        "action",
                        "engineer",
                    )
                    if name in values and values[name] not in (None, "")
                ]
                summary = "; ".join(
                    f"{name.replace('_', ' ').title()}: {values[name]}"
                    for name in useful
                )
            history.append(
                {
                    "id": int(row["id"]),
                    "changed_at": row["changed_at"],
                    "table_name": row["table_name"],
                    "area": area_labels.get(
                        row["table_name"],
                        str(row["table_name"]).replace("_", " ").title(),
                    ),
                    "change_type": row["change_type"],
                    "record_key": row["record_key"],
                    "node": node,
                    "fields": ", ".join(changed_fields),
                    "summary": summary or "Record data changed",
                    "old_values": old_values,
                    "new_values": new_values,
                }
            )
        return history

    def fetch_snapshots(self) -> list[sqlite3.Row]:
        with self.connection() as connection:
            return list(connection.execute("SELECT * FROM snapshots ORDER BY id DESC"))

    def add_floor(self, name: str, level_order: int, dxf_path: str | None = None) -> int:
        with self.connection() as connection:
            cursor = connection.execute(
                "INSERT INTO floors(name, level_order, dxf_path) VALUES (?, ?, ?)",
                (name, level_order, dxf_path),
            )
            return int(cursor.lastrowid)

    def set_floor_dxf(self, floor_id: int, dxf_path: str | None) -> None:
        with self.connection() as connection:
            cursor = connection.execute(
                "UPDATE floors SET dxf_path = ? WHERE id = ?",
                (
                    str(Path(dxf_path).resolve()) if dxf_path else None,
                    int(floor_id),
                ),
            )
            if cursor.rowcount == 0:
                raise ValueError("The selected floor no longer exists.")
            connection.execute(
                "DELETE FROM ignored_zone_shapes WHERE floor_id = ?",
                (int(floor_id),),
            )

    def fetch_floors(self) -> list[sqlite3.Row]:
        with self.connection() as connection:
            return list(connection.execute("SELECT * FROM floors ORDER BY level_order DESC"))

    def assign_zone_geometry(
        self,
        zone: int,
        floor_id: int,
        points: Iterable[tuple[float, float]],
        source_layer: str = "",
    ) -> None:
        geometry_json = json.dumps([[float(x), float(y)] for x, y in points])
        with self.connection() as connection:
            connection.execute(
                """
                UPDATE zones
                SET floor_id = COALESCE(floor_id, ?)
                WHERE number = ?
                """,
                (floor_id, zone),
            )
            connection.execute(
                """
                INSERT INTO zone_geometry(zone, floor_id, geometry_json, source_layer)
                VALUES (?, ?, ?, ?)
                ON CONFLICT(zone, floor_id)
                DO UPDATE SET geometry_json=excluded.geometry_json,
                              source_layer=excluded.source_layer
                """,
                (zone, floor_id, geometry_json, source_layer),
            )

    def fetch_zone_geometry(self) -> list[sqlite3.Row]:
        with self.connection() as connection:
            return list(
                connection.execute(
                    """
                    SELECT g.*, f.name AS floor_name, f.level_order, z.description
                    FROM zone_geometry g
                    JOIN floors f ON f.id = g.floor_id
                    JOIN zones z ON z.number = g.zone
                    ORDER BY f.level_order DESC, g.zone
                    """
                )
            )

    def place_map_asset(
        self,
        entity_kind: str,
        entity_key: str,
        floor_id: int,
        x: float,
        y: float,
        symbol_type: str,
    ) -> None:
        if entity_kind not in {"device", "panel"}:
            raise ValueError(f"Unsupported map entity: {entity_kind}")
        with self.connection() as connection:
            connection.execute(
                """
                INSERT INTO map_assets(entity_kind, entity_key, floor_id, x, y, symbol_type)
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(entity_kind, entity_key) DO UPDATE SET
                    floor_id=excluded.floor_id,
                    x=excluded.x,
                    y=excluded.y,
                    symbol_type=excluded.symbol_type
                """,
                (entity_kind, entity_key, floor_id, float(x), float(y), symbol_type),
            )

    def remove_map_asset(self, entity_kind: str, entity_key: str) -> None:
        with self.connection() as connection:
            connection.execute(
                "DELETE FROM map_assets WHERE entity_kind = ? AND entity_key = ?",
                (entity_kind, entity_key),
            )

    def fetch_map_assets(self, floor_id: int | None = None) -> list[sqlite3.Row]:
        with self.connection() as connection:
            if floor_id is None:
                return list(
                    connection.execute(
                        "SELECT * FROM map_assets ORDER BY floor_id, entity_kind, entity_key"
                    )
                )
            return list(
                connection.execute(
                    """
                    SELECT * FROM map_assets
                    WHERE floor_id = ?
                    ORDER BY entity_kind, entity_key
                    """,
                    (floor_id,),
                )
            )

    def set_output_group_name(self, number: int, name: str) -> None:
        with self.connection() as connection:
            connection.execute(
                """
                INSERT INTO output_groups(number, name) VALUES (?, ?)
                ON CONFLICT(number) DO UPDATE SET name=excluded.name
                """,
                (number, name.strip()),
            )

    def output_group_details(self, number: int) -> tuple[str, list[int]]:
        with self.connection() as connection:
            row = connection.execute(
                "SELECT name FROM output_groups WHERE number = ?", (number,)
            ).fetchone()
            rule_names = [
                item["name"]
                for item in connection.execute(
                    """
                    SELECT DISTINCT name
                    FROM cause_effect_rules
                    WHERE enabled = 1 AND output_group = ?
                    ORDER BY name
                    """,
                    (number,),
                )
            ]
            zones = [
                int(item["trigger_zone"])
                for item in connection.execute(
                    """
                    SELECT DISTINCT trigger_zone
                    FROM cause_effect_rules
                    WHERE enabled = 1 AND output_group = ? AND trigger_zone IS NOT NULL
                    ORDER BY trigger_zone
                    """,
                    (number,),
                )
            ]
        name = row["name"] if row and row["name"] else " / ".join(rule_names)
        return (name, zones)

    def replace_suggested_rules(self, rows: Iterable[tuple]) -> None:
        with self.connection() as connection:
            connection.execute("DELETE FROM cause_effect_rules WHERE source = 'HTM 05-03 Figure 2'")
            connection.executemany(
                """
                INSERT INTO cause_effect_rules(
                    name, trigger_zone, relation, target_zone, target_node,
                    output_group, action, source, enabled, notes
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                list(rows),
            )

    def replace_door_suggested_rules(self, rows: Iterable[tuple]) -> None:
        with self.connection() as connection:
            connection.execute(
                "DELETE FROM cause_effect_rules WHERE source = 'Door drawing suggestion'"
            )
            connection.executemany(
                """
                INSERT INTO cause_effect_rules(
                    name, trigger_zone, relation, target_zone, target_node,
                    output_group, action, source, enabled, notes
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                list(rows),
            )

    def fetch_rules(self) -> list[sqlite3.Row]:
        with self.connection() as connection:
            return list(
                connection.execute(
                    "SELECT * FROM cause_effect_rules ORDER BY trigger_zone, action, target_zone, id"
                )
            )

    def create_door(
        self,
        name: str,
        floor_id: int,
        start: tuple[float, float],
        end: tuple[float, float],
        zone_a: int,
        zone_b: int,
        has_access_control: bool,
        access_device_key: str | None,
        access_normal_state: str,
        has_hold_open: bool,
        hold_open_device_key: str | None,
        hold_open_normal_state: str,
        notes: str = "",
        door_type: str = "SINGLE",
        sprite_position: tuple[float, float] | None = None,
        rotation_degrees: float = 0,
    ) -> int:
        values = self._validated_door_values(
            name,
            floor_id,
            start,
            end,
            zone_a,
            zone_b,
            has_access_control,
            access_device_key,
            access_normal_state,
            has_hold_open,
            hold_open_device_key,
            hold_open_normal_state,
            notes,
            door_type,
            sprite_position,
            rotation_degrees,
        )
        with self.connection() as connection:
            cursor = connection.execute(
                """
                INSERT INTO doors(
                    name, floor_id, start_x, start_y, end_x, end_y,
                    zone_a, zone_b, has_access_control, access_device_key,
                    access_normal_state, has_hold_open, hold_open_device_key,
                    hold_open_normal_state, notes, door_type, sprite_x,
                    sprite_y, rotation_degrees
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                values,
            )
            return int(cursor.lastrowid)

    def update_door(
        self,
        door_id: int,
        name: str,
        floor_id: int,
        start: tuple[float, float],
        end: tuple[float, float],
        zone_a: int,
        zone_b: int,
        has_access_control: bool,
        access_device_key: str | None,
        access_normal_state: str,
        has_hold_open: bool,
        hold_open_device_key: str | None,
        hold_open_normal_state: str,
        notes: str = "",
        door_type: str = "SINGLE",
        sprite_position: tuple[float, float] | None = None,
        rotation_degrees: float = 0,
    ) -> None:
        values = self._validated_door_values(
            name,
            floor_id,
            start,
            end,
            zone_a,
            zone_b,
            has_access_control,
            access_device_key,
            access_normal_state,
            has_hold_open,
            hold_open_device_key,
            hold_open_normal_state,
            notes,
            door_type,
            sprite_position,
            rotation_degrees,
        )
        with self.connection() as connection:
            cursor = connection.execute(
                """
                UPDATE doors
                SET name = ?, floor_id = ?,
                    start_x = ?, start_y = ?, end_x = ?, end_y = ?,
                    zone_a = ?, zone_b = ?,
                    has_access_control = ?, access_device_key = ?,
                    access_normal_state = ?, has_hold_open = ?,
                    hold_open_device_key = ?, hold_open_normal_state = ?,
                    notes = ?, door_type = ?, sprite_x = ?, sprite_y = ?,
                    rotation_degrees = ?
                WHERE id = ?
                """,
                (*values, int(door_id)),
            )
            if cursor.rowcount == 0:
                raise ValueError("The selected door no longer exists.")

    def move_door(
        self,
        door_id: int,
        sprite_x: float,
        sprite_y: float,
    ) -> None:
        with self.connection() as connection:
            cursor = connection.execute(
                """
                UPDATE doors
                SET sprite_x = ?, sprite_y = ?
                WHERE id = ?
                """,
                (float(sprite_x), float(sprite_y), int(door_id)),
            )
            if cursor.rowcount == 0:
                raise ValueError("The selected door no longer exists.")

    def _validated_door_values(
        self,
        name: str,
        floor_id: int,
        start: tuple[float, float],
        end: tuple[float, float],
        zone_a: int,
        zone_b: int,
        has_access_control: bool,
        access_device_key: str | None,
        access_normal_state: str,
        has_hold_open: bool,
        hold_open_device_key: str | None,
        hold_open_normal_state: str,
        notes: str,
        door_type: str = "SINGLE",
        sprite_position: tuple[float, float] | None = None,
        rotation_degrees: float = 0,
    ) -> tuple:
        name = name.strip()
        access_device_key = (
            str(access_device_key).strip() if access_device_key else None
        )
        hold_open_device_key = (
            str(hold_open_device_key).strip() if hold_open_device_key else None
        )
        access_normal_state = access_normal_state.strip().upper()
        hold_open_normal_state = hold_open_normal_state.strip().upper()
        door_type = door_type.strip().upper()
        if not name:
            raise ValueError("Enter a door name.")
        if not has_access_control and not has_hold_open:
            raise ValueError("Select access control, fire hold-open, or both.")
        if access_normal_state not in {"LOCKED", "UNLOCKED"}:
            raise ValueError("Access state must be LOCKED or UNLOCKED.")
        if hold_open_normal_state not in {"HELD OPEN", "CLOSED"}:
            raise ValueError("Hold-open state must be HELD OPEN or CLOSED.")
        if door_type not in {"SINGLE", "DOUBLE"}:
            raise ValueError("Door type must be SINGLE or DOUBLE.")
        start_x, start_y = map(float, start)
        end_x, end_y = map(float, end)
        if math.hypot(end_x - start_x, end_y - start_y) < 0.001:
            raise ValueError("Draw the door between two different points.")
        with self.connection() as connection:
            floor_exists = connection.execute(
                "SELECT 1 FROM floors WHERE id = ?", (int(floor_id),)
            ).fetchone()
            zones = {
                int(row["number"])
                for row in connection.execute(
                    "SELECT number FROM zones WHERE number IN (?, ?)",
                    (int(zone_a), int(zone_b)),
                )
            }
            snapshot_id = self.latest_snapshot_id()
            linked_keys = {
                key
                for key in (access_device_key, hold_open_device_key)
                if key is not None
            }
            existing_keys = (
                {
                    str(row["stable_key"])
                    for row in connection.execute(
                        f"""
                        SELECT stable_key FROM devices
                        WHERE snapshot_id = ?
                          AND stable_key IN ({','.join('?' for _ in linked_keys)})
                        """,
                        (snapshot_id, *sorted(linked_keys)),
                    )
                }
                if snapshot_id is not None and linked_keys
                else set()
            )
        if floor_exists is None:
            raise ValueError("The selected floor no longer exists.")
        if zones != {int(zone_a), int(zone_b)}:
            raise ValueError("One or both selected door zones no longer exist.")
        if linked_keys != existing_keys:
            raise ValueError("One or both linked fire-alarm devices no longer exist.")
        if sprite_position is None:
            sprite_x = (start_x + end_x) / 2.0
            sprite_y = (start_y + end_y) / 2.0
        else:
            sprite_x, sprite_y = map(float, sprite_position)
        return (
            name,
            int(floor_id),
            start_x,
            start_y,
            end_x,
            end_y,
            int(zone_a),
            int(zone_b),
            int(bool(has_access_control)),
            access_device_key if has_access_control else None,
            access_normal_state,
            int(bool(has_hold_open)),
            hold_open_device_key if has_hold_open else None,
            hold_open_normal_state,
            notes.strip(),
            door_type,
            sprite_x,
            sprite_y,
            float(rotation_degrees),
        )

    def fetch_doors(self, floor_id: int | None = None) -> list[sqlite3.Row]:
        with self.connection() as connection:
            if floor_id is None:
                return list(
                    connection.execute(
                        """
                        SELECT d.*, f.name AS floor_name
                        FROM doors d
                        JOIN floors f ON f.id = d.floor_id
                        ORDER BY f.level_order DESC, d.name, d.id
                        """
                    )
                )
            return list(
                connection.execute(
                    """
                    SELECT d.*, f.name AS floor_name
                    FROM doors d
                    JOIN floors f ON f.id = d.floor_id
                    WHERE d.floor_id = ?
                    ORDER BY d.name, d.id
                    """,
                    (int(floor_id),),
                )
            )

    def delete_door(self, door_id: int) -> None:
        with self.connection() as connection:
            connection.execute("DELETE FROM doors WHERE id = ?", (int(door_id),))

    def suggest_door_control_devices(
        self,
        zone_a: int,
        zone_b: int,
        capability: str,
    ) -> list[sqlite3.Row]:
        capability = capability.strip().casefold()
        keywords = (
            ("access", "door", "lock", "release", "relay")
            if capability == "access"
            else ("hold", "door", "magnet", "release", "relay")
        )

        def score(row: sqlite3.Row) -> tuple:
            text = " ".join(
                str(row[key] or "")
                for key in ("text", "observed_type", "output_group_name")
            ).casefold()
            is_output = (
                (
                    row["output_group"] is not None
                    and int(row["output_group"]) > 0
                )
                or "output" in text
                or "relay" in text
            )
            keyword_matches = sum(
                keyword in text for keyword in keywords
            )
            return (
                int(not is_output),
                int(int(row["zone"]) not in {int(zone_a), int(zone_b)}),
                -keyword_matches,
                int(row["node"]),
                int(row["loop"]),
                int(row["address"]),
                int(row["sub_address"]),
            )

        return sorted(self.fetch_devices(), key=score)

    def update_zone_geometry(
        self,
        geometry_id: int,
        points: Iterable[tuple[float, float]],
    ) -> None:
        geometry_json = json.dumps([[float(x), float(y)] for x, y in points])
        with self.connection() as connection:
            connection.execute(
                """
                UPDATE zone_geometry
                SET geometry_json = ?
                WHERE id = ?
                """,
                (geometry_json, int(geometry_id)),
            )

    def reassign_zone_geometry(self, geometry_id: int, zone: int) -> None:
        with self.connection() as connection:
            row = connection.execute(
                "SELECT zone, floor_id FROM zone_geometry WHERE id = ?",
                (int(geometry_id),),
            ).fetchone()
            if row is None:
                raise ValueError("The selected zone geometry no longer exists.")
            existing = connection.execute(
                """
                SELECT id FROM zone_geometry
                WHERE zone = ? AND floor_id = ? AND id <> ?
                """,
                (int(zone), int(row["floor_id"]), int(geometry_id)),
            ).fetchone()
            if existing is not None:
                raise ValueError(
                    "That zone already has an assigned polygon on this floor."
                )
            connection.execute(
                "UPDATE zone_geometry SET zone = ? WHERE id = ?",
                (int(zone), int(geometry_id)),
            )
            connection.execute(
                """
                UPDATE zones
                SET floor_id = COALESCE(floor_id, ?)
                WHERE number = ?
                """,
                (int(row["floor_id"]), int(zone)),
            )
            old_zone_remaining = connection.execute(
                "SELECT floor_id FROM zone_geometry WHERE zone = ? LIMIT 1",
                (int(row["zone"]),),
            ).fetchone()
            connection.execute(
                "UPDATE zones SET floor_id = ? WHERE number = ?",
                (
                    (
                        int(old_zone_remaining["floor_id"])
                        if old_zone_remaining is not None
                        else None
                    ),
                    int(row["zone"]),
                ),
            )

    def remove_zone_geometry(self, geometry_id: int) -> None:
        with self.connection() as connection:
            row = connection.execute(
                "SELECT zone, floor_id FROM zone_geometry WHERE id = ?",
                (int(geometry_id),),
            ).fetchone()
            if row is None:
                return
            connection.execute(
                "DELETE FROM zone_geometry WHERE id = ?",
                (int(geometry_id),),
            )
            remaining = connection.execute(
                """
                SELECT floor_id
                FROM zone_geometry
                WHERE zone = ?
                ORDER BY floor_id
                LIMIT 1
                """,
                (int(row["zone"]),),
            ).fetchone()
            connection.execute(
                "UPDATE zones SET floor_id = ? WHERE number = ?",
                (
                    int(remaining["floor_id"]) if remaining is not None else None,
                    int(row["zone"]),
                ),
            )

    def ignore_zone_shape(
        self,
        floor_id: int,
        points: Iterable[tuple[float, float]],
        source_layer: str = "",
        shape_key: str | None = None,
    ) -> None:
        point_list = [
            (float(x), float(y))
            for x, y in points
        ]
        key = shape_key or zone_shape_key(point_list)
        with self.connection() as connection:
            connection.execute(
                """
                INSERT OR REPLACE INTO ignored_zone_shapes(
                    floor_id, shape_key, geometry_json, source_layer
                ) VALUES (?, ?, ?, ?)
                """,
                (
                    int(floor_id),
                    key,
                    json.dumps([[x, y] for x, y in point_list]),
                    str(source_layer or ""),
                ),
            )

    def fetch_ignored_zone_shape_keys(self, floor_id: int) -> set[str]:
        with self.connection() as connection:
            return {
                str(row["shape_key"])
                for row in connection.execute(
                    """
                    SELECT shape_key
                    FROM ignored_zone_shapes
                    WHERE floor_id = ?
                    """,
                    (int(floor_id),),
                )
            }

    def import_cause_effect(
        self, workbook_path: str | Path
    ) -> tuple[int, CauseEffectWorkbook]:
        parsed = read_cause_effect_workbook(workbook_path)
        with self.connection() as connection:
            existing = connection.execute(
                "SELECT id FROM cause_effect_imports WHERE sha256 = ?",
                (parsed.sha256,),
            ).fetchone()
            if existing:
                return int(existing["id"]), parsed

            previous_comments: dict[tuple[str, int, int, str], str] = {}
            latest = connection.execute(
                "SELECT MAX(id) AS id FROM cause_effect_imports"
            ).fetchone()["id"]
            if latest is not None:
                for row in connection.execute(
                    """
                    SELECT trigger_zone, target_node, output_group, ringing_style,
                           comments
                    FROM cause_effect_activations
                    WHERE import_id = ? AND TRIM(comments) <> ''
                    """,
                    (latest,),
                ):
                    previous_comments[
                        (
                            row["trigger_zone"],
                            row["target_node"],
                            row["output_group"],
                            row["ringing_style"],
                        )
                    ] = row["comments"]

            cursor = connection.execute(
                """
                INSERT INTO cause_effect_imports(
                    imported_at, source_name, source_path, sha256,
                    activation_count, reference_count, matched_count,
                    matrix_only_count, reference_only_count,
                    activation_codes_json, warnings_json
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    _now(),
                    parsed.source.name,
                    str(parsed.source),
                    parsed.sha256,
                    len(parsed.activations),
                    parsed.reference_count,
                    parsed.matched_count,
                    parsed.matrix_only_count,
                    parsed.reference_only_count,
                    json.dumps(parsed.activation_codes),
                    json.dumps(parsed.warnings),
                ),
            )
            import_id = int(cursor.lastrowid)
            self._insert_cause_effect_output_groups(
                connection,
                import_id,
                parsed,
            )
            connection.executemany(
                """
                INSERT INTO cause_effect_activations(
                    import_id, trigger_zone, trigger_zone_name,
                    target_node, target_node_name, output_group,
                    output_group_name, output_zone_name, ringing_style,
                    source_row, source_column, reference_status, comments
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    (
                        import_id,
                        activation.trigger_zone,
                        activation.trigger_zone_name,
                        activation.target_node,
                        activation.target_node_name,
                        activation.output_group,
                        activation.output_group_name,
                        activation.output_zone_name,
                        activation.activation_code,
                        activation.source_row,
                        activation.source_column,
                        activation.reference_status,
                        previous_comments.get(
                            (
                                activation.trigger_zone,
                                activation.target_node,
                                activation.output_group,
                                activation.activation_code,
                            ),
                            "",
                        ),
                    )
                    for activation in parsed.activations
                ],
            )
            connection.executemany(
                """
                INSERT INTO cause_effect_reference_only(
                    import_id, trigger_zone, target_node, target_node_name,
                    output_group, output_group_name, output_zone_name,
                    ringing_style, source_row
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    (
                        import_id,
                        reference.trigger_zone,
                        reference.target_node,
                        reference.target_node_name,
                        reference.output_group,
                        reference.output_group_name,
                        reference.output_zone_name,
                        reference.activation_code,
                        reference.source_row,
                    )
                    for reference in parsed.reference_only
                ],
            )
            changes = self._calculate_cause_effect_changes(
                connection,
                latest,
                import_id,
            )
            connection.executemany(
                """
                INSERT INTO cause_effect_changes(
                    import_id, entity, stable_key, change_type,
                    field, old_value, new_value
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    (
                        import_id,
                        change.entity,
                        change.stable_key,
                        change.change_type,
                        change.field,
                        change.old_value,
                        change.new_value,
                    )
                    for change in changes
                ],
            )
        return import_id, parsed

    @staticmethod
    def _insert_cause_effect_output_groups(
        connection: sqlite3.Connection,
        import_id: int,
        parsed: CauseEffectWorkbook,
    ) -> None:
        connection.executemany(
            """
            INSERT OR REPLACE INTO cause_effect_output_groups(
                import_id, target_node, target_node_name,
                output_group, output_group_name, source_column
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    import_id,
                    output.target_node,
                    output.target_node_name,
                    output.output_group,
                    output.output_group_name,
                    get_column_letter(output.index + 1),
                )
                for output in parsed.output_groups
            ],
        )

    @staticmethod
    def _calculate_cause_effect_changes(
        connection: sqlite3.Connection,
        previous_import: int | None,
        import_id: int,
    ) -> list[Change]:
        if previous_import is None:
            count = connection.execute(
                """
                SELECT COUNT(*) AS count
                FROM cause_effect_activations
                WHERE import_id = ?
                """,
                (import_id,),
            ).fetchone()["count"]
            return [
                Change(
                    entity="cause_effect_matrix",
                    stable_key=str(import_id),
                    change_type="initial",
                    field="activations",
                    old_value=None,
                    new_value=str(count),
                )
            ]

        old_rows = {
            _cause_effect_stable_key(row): row
            for row in connection.execute(
                """
                SELECT * FROM cause_effect_activations
                WHERE import_id = ?
                """,
                (previous_import,),
            )
        }
        new_rows = {
            _cause_effect_stable_key(row): row
            for row in connection.execute(
                """
                SELECT * FROM cause_effect_activations
                WHERE import_id = ?
                """,
                (import_id,),
            )
        }
        old_outputs = {
            _cause_effect_output_key(row): row
            for row in connection.execute(
                """
                SELECT * FROM cause_effect_output_groups
                WHERE import_id = ?
                """,
                (previous_import,),
            )
        }
        new_outputs = {
            _cause_effect_output_key(row): row
            for row in connection.execute(
                """
                SELECT * FROM cause_effect_output_groups
                WHERE import_id = ?
                """,
                (import_id,),
            )
        }
        changes: list[Change] = []
        for key in sorted(new_rows.keys() - old_rows.keys()):
            changes.append(
                Change(
                    "cause_effect_activation",
                    key,
                    "added",
                    None,
                    None,
                    _cause_effect_activation_summary(new_rows[key]),
                )
            )
        for key in sorted(old_rows.keys() - new_rows.keys()):
            changes.append(
                Change(
                    "cause_effect_activation",
                    key,
                    "removed",
                    None,
                    _cause_effect_activation_summary(old_rows[key]),
                    None,
                )
            )
        for key in sorted(old_rows.keys() & new_rows.keys()):
            old_value = old_rows[key]["ringing_style"]
            new_value = new_rows[key]["ringing_style"]
            if old_value != new_value:
                changes.append(
                    Change(
                        "cause_effect_activation",
                        key,
                        "modified",
                        "ringing_style",
                        _stringify(old_value),
                        _stringify(new_value),
                    )
                )

        for key in sorted(new_outputs.keys() - old_outputs.keys()):
            changes.append(
                Change(
                    "cause_effect_output_group",
                    key,
                    "added",
                    None,
                    None,
                    _cause_effect_output_summary(new_outputs[key]),
                )
            )
        for key in sorted(old_outputs.keys() - new_outputs.keys()):
            changes.append(
                Change(
                    "cause_effect_output_group",
                    key,
                    "removed",
                    None,
                    _cause_effect_output_summary(old_outputs[key]),
                    None,
                )
            )
        for key in sorted(old_outputs.keys() & new_outputs.keys()):
            old_name = old_outputs[key]["output_group_name"]
            new_name = new_outputs[key]["output_group_name"]
            if old_name != new_name:
                changes.append(
                    Change(
                        "cause_effect_output_group",
                        key,
                        "modified",
                        "output_group_name",
                        _stringify(old_name),
                        _stringify(new_name),
                    )
                )

        label_specs = (
            (
                "cause_effect_zone",
                "trigger_zone_name",
                old_rows.values(),
                new_rows.values(),
                lambda row: f"zone {row['trigger_zone']}",
            ),
            (
                "cause_effect_node",
                "target_node_name",
                old_outputs.values(),
                new_outputs.values(),
                lambda row: f"node {row['target_node']}",
            ),
        )
        for entity, field, old_values, new_values, key_for_row in label_specs:
            old_labels = {key_for_row(row): row[field] for row in old_values}
            new_labels = {key_for_row(row): row[field] for row in new_values}
            for key in sorted(old_labels.keys() & new_labels.keys()):
                if old_labels[key] != new_labels[key]:
                    changes.append(
                        Change(
                            entity,
                            key,
                            "modified",
                            field,
                            _stringify(old_labels[key]),
                            _stringify(new_labels[key]),
                        )
                    )
        return changes

    def fetch_cause_effect_changes(
        self,
        import_id: int | None = None,
    ) -> list[sqlite3.Row]:
        with self.connection() as connection:
            if import_id is None:
                import_id = connection.execute(
                    "SELECT MAX(id) AS id FROM cause_effect_imports"
                ).fetchone()["id"]
            if import_id is None:
                return []
            return list(
                connection.execute(
                    """
                    SELECT * FROM cause_effect_changes
                    WHERE import_id = ?
                    ORDER BY id
                    """,
                    (import_id,),
                )
            )

    def latest_cause_effect_import(self) -> sqlite3.Row | None:
        with self.connection() as connection:
            return connection.execute(
                "SELECT * FROM cause_effect_imports ORDER BY id DESC LIMIT 1"
            ).fetchone()

    def fetch_cause_effect_output_groups(self) -> list[sqlite3.Row]:
        latest = self.latest_cause_effect_import()
        if latest is None:
            return []
        with self.connection() as connection:
            return list(
                connection.execute(
                    """
                    SELECT * FROM cause_effect_output_groups
                    WHERE import_id = ?
                    ORDER BY target_node, output_group
                    """,
                    (latest["id"],),
                )
            )

    def fetch_cause_effect_activations(
        self,
        trigger_zone: object | None = None,
        scope_node: int | None = None,
    ) -> list:
        latest = self.latest_cause_effect_import()
        if latest is None:
            return self._fetch_configuration_activations(
                trigger_zone,
                scope_node,
            )
        clauses = ["import_id = ?"]
        parameters: list[object] = [latest["id"]]
        if trigger_zone is not None:
            clauses.append("trigger_zone = ?")
            parameters.append(normalise_zone_key(trigger_zone))
        if scope_node is not None:
            clauses.append("target_node = ?")
            parameters.append(int(scope_node))
        with self.connection() as connection:
            return list(
                connection.execute(
                    f"""
                    SELECT * FROM cause_effect_activations
                    WHERE {' AND '.join(clauses)}
                    ORDER BY CASE
                                 WHEN trigger_zone GLOB '[0-9]*' THEN 0
                                 ELSE 1
                             END,
                             CAST(trigger_zone AS REAL), trigger_zone,
                             target_node, output_group, id
                    """,
                    parameters,
                )
            )

    def fetch_configuration_output_group_lines(
        self,
        scope_node: int | None = None,
    ) -> list[sqlite3.Row]:
        snapshot_id = self.latest_snapshot_id()
        if snapshot_id is None:
            return []
        clauses = ["snapshot_id = ?"]
        parameters: list[object] = [snapshot_id]
        if scope_node is not None:
            clauses.append("target_node = ?")
            parameters.append(int(scope_node))
        with self.connection() as connection:
            return list(
                connection.execute(
                    f"""
                    SELECT *
                    FROM configuration_output_group_lines
                    WHERE {' AND '.join(clauses)}
                    ORDER BY target_node, output_group, source_row
                    """,
                    parameters,
                )
            )

    def _fetch_configuration_activations(
        self,
        trigger_zone: object | None,
        scope_node: int | None,
    ) -> list[dict]:
        if trigger_zone is None:
            return []
        normalised = normalise_zone_key(trigger_zone)
        try:
            numeric_zone = float(normalised)
        except ValueError:
            return []
        if not numeric_zone.is_integer() or numeric_zone <= 0:
            return []
        zone = int(numeric_zone)
        activations = []
        for line in self.fetch_configuration_output_group_lines(scope_node):
            if (
                int(line["operation"]) != 0
                or zone < int(line["zone_from"])
                or zone > int(line["zone_to"])
            ):
                continue
            activations.append(
                {
                    "id": (
                        f"skf/{line['source_row']}/"
                        f"{line['target_node']}/{line['output_group']}"
                    ),
                    "trigger_zone": str(zone),
                    "trigger_zone_name": "",
                    "target_node": int(line["target_node"]),
                    "target_node_name": str(line["target_node_name"] or ""),
                    "output_group": int(line["output_group"]),
                    "output_group_name": str(
                        line["output_group_name"] or ""
                    ),
                    "output_zone_name": "",
                    "ringing_style": str(line["ringing_style"] or "Triggered"),
                    "ringing_style_name": str(
                        line["ringing_style_name"] or ""
                    ),
                    "zone_qualifiers": int(line["zone_qualifiers"]),
                    "source_row": int(line["source_row"]),
                    "comments": "",
                    "reference_status": "configuration",
                }
            )
        return activations

    def fetch_cause_effect_trigger_zones(self) -> list[sqlite3.Row]:
        latest = self.latest_cause_effect_import()
        if latest is None:
            return []
        with self.connection() as connection:
            return list(
                connection.execute(
                    """
                    SELECT trigger_zone,
                           MAX(trigger_zone_name) AS trigger_zone_name,
                           COUNT(*) AS activation_count
                    FROM cause_effect_activations
                    WHERE import_id = ?
                    GROUP BY trigger_zone
                    ORDER BY CASE
                                 WHEN trigger_zone GLOB '[0-9]*' THEN 0
                                 ELSE 1
                             END,
                             CAST(trigger_zone AS REAL), trigger_zone
                    """,
                    (latest["id"],),
                )
            )

    def fetch_cause_effect_reference_only(self) -> list[sqlite3.Row]:
        latest = self.latest_cause_effect_import()
        if latest is None:
            return []
        with self.connection() as connection:
            return list(
                connection.execute(
                    """
                    SELECT * FROM cause_effect_reference_only
                    WHERE import_id = ?
                    ORDER BY source_row, id
                    """,
                    (latest["id"],),
                )
            )

    def update_cause_effect_comment(
        self,
        activation_id: int,
        comments: str,
    ) -> None:
        latest = self.latest_cause_effect_import()
        if latest is None:
            return
        with self.connection() as connection:
            connection.execute(
                """
                UPDATE cause_effect_activations
                SET comments = ?
                WHERE id = ? AND import_id = ?
                """,
                (comments.strip(), int(activation_id), latest["id"]),
            )

    def add_rule(
        self,
        name: str,
        trigger_zone: int,
        relation: str,
        target_zone: int | None,
        target_node: int | None,
        output_group: int | None,
        action: str,
        notes: str = "",
    ) -> int:
        with self.connection() as connection:
            cursor = connection.execute(
                """
                INSERT INTO cause_effect_rules(
                    name, trigger_zone, relation, target_zone, target_node,
                    output_group, action, source, enabled, notes
                ) VALUES (?, ?, ?, ?, ?, ?, ?, 'custom', 1, ?)
                """,
                (
                    name,
                    trigger_zone,
                    relation,
                    target_zone,
                    target_node,
                    output_group,
                    action,
                    notes,
                ),
            )
            return int(cursor.lastrowid)

    def fetch_rule(self, rule_id: int) -> sqlite3.Row | None:
        with self.connection() as connection:
            return connection.execute(
                "SELECT * FROM cause_effect_rules WHERE id = ?",
                (int(rule_id),),
            ).fetchone()

    def update_custom_rule(
        self,
        rule_id: int,
        name: str,
        trigger_zone: int,
        relation: str,
        target_zone: int | None,
        target_node: int | None,
        output_group: int | None,
        action: str,
        notes: str = "",
    ) -> None:
        with self.connection() as connection:
            cursor = connection.execute(
                """
                UPDATE cause_effect_rules
                SET name = ?, trigger_zone = ?, relation = ?,
                    target_zone = ?, target_node = ?, output_group = ?,
                    action = ?, notes = ?
                WHERE id = ? AND source = 'custom'
                """,
                (
                    name,
                    trigger_zone,
                    relation,
                    target_zone,
                    target_node,
                    output_group,
                    action,
                    notes,
                    int(rule_id),
                ),
            )
            if cursor.rowcount != 1:
                raise ValueError(
                    "Only custom cause-and-effect rules can be edited."
                )

    def delete_custom_rule(self, rule_id: int) -> None:
        self.delete_custom_rules([rule_id])

    def delete_custom_rules(self, rule_ids: Iterable[int]) -> None:
        ids = sorted({int(rule_id) for rule_id in rule_ids})
        if not ids:
            return
        with self.connection() as connection:
            placeholders = ", ".join("?" for _rule_id in ids)
            custom_count = connection.execute(
                f"""
                SELECT COUNT(*)
                FROM cause_effect_rules
                WHERE id IN ({placeholders}) AND source = 'custom'
                """,
                ids,
            ).fetchone()[0]
            if int(custom_count) != len(ids):
                raise ValueError(
                    "Only custom cause-and-effect rules can be removed "
                    "through this method."
                )
        self.delete_rules(ids)

    def delete_rules(self, rule_ids: Iterable[int]) -> None:
        ids = sorted({int(rule_id) for rule_id in rule_ids})
        if not ids:
            return
        placeholders = ", ".join("?" for _rule_id in ids)
        with self.connection() as connection:
            existing_count = connection.execute(
                f"""
                SELECT COUNT(*)
                FROM cause_effect_rules
                WHERE id IN ({placeholders})
                """,
                ids,
            ).fetchone()[0]
            if int(existing_count) != len(ids):
                raise ValueError("One or more selected rules no longer exist.")
            connection.execute(
                f"""
                DELETE FROM cause_effect_rules
                WHERE id IN ({placeholders})
                """,
                ids,
            )

    def set_node_power(
        self,
        node: int,
        battery_ah: float | None,
        standby_hours: float,
        alarm_minutes: float,
        safety_factor: float = 1.25,
    ) -> None:
        with self.connection() as connection:
            connection.execute(
                """
                INSERT INTO node_power(node, battery_ah, standby_hours, alarm_minutes, safety_factor)
                VALUES (?, ?, ?, ?, ?)
                ON CONFLICT(node) DO UPDATE SET
                    battery_ah=excluded.battery_ah,
                    standby_hours=excluded.standby_hours,
                    alarm_minutes=excluded.alarm_minutes,
                    safety_factor=excluded.safety_factor
                """,
                (node, battery_ah, standby_hours, alarm_minutes, safety_factor),
            )

    def fetch_node_power(self) -> dict[int, sqlite3.Row]:
        with self.connection() as connection:
            return {
                int(row["node"]): row
                for row in connection.execute("SELECT * FROM node_power")
            }

    def create_test_session(
        self,
        engineer: str,
        scope_node: int | None,
        trigger_zone: int | float | str,
        results: Iterable[tuple[str | None, str, str, str]],
        notes: str = "",
    ) -> int:
        """Persist the current simulation/check as a commissioning test session."""
        with self.connection() as connection:
            cursor = connection.execute(
                """
                INSERT INTO test_sessions(started_at, completed_at, engineer, scope_node, notes)
                VALUES (?, ?, ?, ?, ?)
                """,
                (_now(), _now(), engineer, scope_node, notes),
            )
            session_id = int(cursor.lastrowid)
            connection.executemany(
                """
                INSERT INTO test_results(
                    session_id, trigger_zone, stable_key, expected_state,
                    actual_state, result, comments, tested_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    (
                        session_id,
                        trigger_zone,
                        stable_key,
                        expected_state,
                        result,
                        _normalise_test_result(result),
                        comments,
                        _now(),
                    )
                    for stable_key, expected_state, result, comments in results
                ],
            )
            return session_id

    def fetch_test_sessions(self) -> list[sqlite3.Row]:
        with self.connection() as connection:
            return list(
                connection.execute(
                    """
                    SELECT s.*, COUNT(r.id) AS result_count
                    FROM test_sessions s
                    LEFT JOIN test_results r ON r.session_id = s.id
                    GROUP BY s.id
                    ORDER BY s.id DESC
                    """
                )
            )

    def import_test_sessions(self, sessions) -> tuple[int, int]:
        """Append spreadsheet test sessions and their output-group results."""
        session_count = 0
        result_count = 0
        with self.connection() as connection:
            for session in sessions:
                notes = str(session.notes or "").strip()
                audit_note = f"Imported workbook session: {session.session_key}"
                notes = f"{notes}\n{audit_note}".strip()
                cursor = connection.execute(
                    """
                    INSERT INTO test_sessions(
                        started_at, completed_at, engineer, scope_node, notes
                    ) VALUES (?, ?, ?, ?, ?)
                    """,
                    (
                        _now(),
                        _now(),
                        session.engineer,
                        session.scope_node,
                        notes,
                    ),
                )
                session_id = int(cursor.lastrowid)
                rows = [
                    (
                        session_id,
                        session.trigger_zone,
                        stable_key,
                        expected_state,
                        actual_state,
                        result,
                        comments,
                        tested_at or _now(),
                    )
                    for (
                        stable_key,
                        expected_state,
                        actual_state,
                        result,
                        comments,
                        tested_at,
                    ) in session.results
                ]
                connection.executemany(
                    """
                    INSERT INTO test_results(
                        session_id, trigger_zone, stable_key, expected_state,
                        actual_state, result, comments, tested_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    rows,
                )
                session_count += 1
                result_count += len(rows)
        return session_count, result_count

    def fetch_test_results(self, session_id: int | None = None) -> list[sqlite3.Row]:
        with self.connection() as connection:
            if session_id is None:
                return list(
                    connection.execute(
                        "SELECT * FROM test_results ORDER BY session_id, id"
                    )
                )
            return list(
                connection.execute(
                    """
                    SELECT * FROM test_results
                    WHERE session_id = ?
                    ORDER BY id
                    """,
                    (int(session_id),),
                )
            )


def _configuration_output_group_values(
    connection: sqlite3.Connection,
    snapshot_id: int,
) -> dict[str, dict]:
    grouped: dict[tuple[int, int], list[sqlite3.Row]] = {}
    for row in connection.execute(
        """
        SELECT *
        FROM configuration_output_group_lines
        WHERE snapshot_id = ?
        ORDER BY target_node, output_group, source_row
        """,
        (snapshot_id,),
    ):
        grouped.setdefault(
            (int(row["target_node"]), int(row["output_group"])),
            [],
        ).append(row)
    values: dict[str, dict] = {}
    for (node, output_group), rows in grouped.items():
        def combined(field: str) -> str:
            return " / ".join(
                dict.fromkeys(
                    str(row[field] or "").strip()
                    for row in rows
                    if str(row[field] or "").strip()
                )
            )

        zone_triggers = "; ".join(
            dict.fromkeys(
                (
                    f"{int(row['zone_from'])}"
                    if int(row["zone_from"]) == int(row["zone_to"])
                    else (
                        f"{int(row['zone_from'])}-"
                        f"{int(row['zone_to'])}"
                    )
                )
                for row in rows
            )
        )
        ringing_styles = "; ".join(
            dict.fromkeys(
                " ".join(
                    part
                    for part in (
                        f"Style {int(row['output_style_number'])}",
                        str(row["ringing_style"] or "").strip(),
                        str(row["ringing_style_name"] or "").strip(),
                    )
                    if part
                )
                for row in rows
            )
        )
        operations = "; ".join(
            dict.fromkeys(str(int(row["operation"])) for row in rows)
        )
        qualifiers = "; ".join(
            dict.fromkeys(
                str(int(row["zone_qualifiers"])) for row in rows
            )
        )
        key = f"node {node}/group {output_group}"
        values[key] = {
            "node": node,
            "target_node_name": combined("target_node_name"),
            "output_group": output_group,
            "output_group_name": combined("output_group_name"),
            "zone_triggers": zone_triggers,
            "ringing_styles": ringing_styles,
            "operations": operations,
            "zone_qualifiers": qualifiers,
        }
    return values


def _output_group_change_summary(group: dict) -> str:
    name = str(group.get("output_group_name") or "").strip()
    return (
        f"Group {group['output_group']}"
        f"{' - ' + name if name else ''}; "
        f"zones {group['zone_triggers']}; "
        f"{group['ringing_styles']}"
    )


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _stringify(value: object) -> str | None:
    return None if value is None else str(value)


def _cause_effect_stable_key(row: sqlite3.Row) -> str:
    return (
        f"zone {row['trigger_zone']} -> "
        f"node {row['target_node']}/group {row['output_group']}"
    )


def _cause_effect_activation_summary(row: sqlite3.Row) -> str:
    ringing_style = str(row["ringing_style"] or "Not specified")
    output_name = str(row["output_group_name"] or "").strip()
    return (
        f"{ringing_style}"
        f"{' - ' + output_name if output_name else ''}"
    )


def _cause_effect_output_key(row: sqlite3.Row) -> str:
    return f"node {row['target_node']}/group {row['output_group']}"


def _cause_effect_output_summary(row: sqlite3.Row) -> str:
    return str(row["output_group_name"] or "Unnamed output group")


def _normalise_test_result(value: str) -> str:
    normalised = value.strip().casefold()
    if normalised in {"pass", "passed", "ok", "correct"}:
        return "pass"
    if normalised in {"fail", "failed", "incorrect"}:
        return "fail"
    return "not-tested" if not normalised else "observation"
